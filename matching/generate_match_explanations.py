#!/usr/bin/env python3
"""
Generate LLM-based explanations for matched pairs and export the results.

Usage examples:
    python generate_match_explanations.py --mode heuristic --limit 5
    python generate_match_explanations.py --model gpt-4.1-mini
    python generate_match_explanations.py --api-key sk-... --api-base https://api.openai.com/v1
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import time
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd
from dotenv import load_dotenv
from openai import OpenAI

from matching.models import User
from matching.parser import load_users
from matching.scoring import compute_score_breakdown


load_dotenv()


@dataclass
class QuestionMeta:
    number: int
    title: str
    section: str = ""
    question_type: str = ""
    options: list[str] = field(default_factory=list)
    min_selections: int | None = None
    max_selections: int | None = None
    left_label: str = ""
    right_label: str = ""

QUESTION_GROUPS: list[tuple[str, range]] = [
    ("基础匹配条件", range(1, 13)),
    ("性格与吸引", range(13, 21)),
    ("生活方式", range(21, 32)),
    ("兴趣偏好", range(32, 34)),
    ("依恋与亲密关系", range(34, 44)),
    ("长期关系与未来", range(44, 53)),
]

PROFILE_QUESTIONS: list[int] = [
    3, 5, 6, 7, 8, 9, 10, 11, 12,
    17, 18, 20,
    23, 24, 25, 26, 27, 28, 29, 30, 31,
    32, 33,
    34, 35, 36, 37, 38, 39, 40, 41, 42, 43,
    44, 45, 46, 47, 48, 49, 50, 51, 52,
]

DIMENSION_LABELS = {
    "D1_purpose_stage": "匹配目的与阶段",
    "D2_geography": "地域与未来规划",
    "D3_attraction": "吸引力与性格互补",
    "D4_lifestyle": "生活方式",
    "D5_interests": "兴趣重合",
    "D6_attachment": "安全感与亲密关系",
    "D7_future": "长期关系与未来观",
}

QUESTION_HEADING_RE = re.compile(r"^###\s+(\d+)\.\s+(.*)$")


def parse_choice_limit(value: str) -> tuple[int | None, int | None]:
    numbers = [int(item) for item in re.findall(r"\d+", value)]
    if not numbers:
        return None, None
    if len(numbers) == 1:
        return numbers[0], numbers[0]
    return numbers[0], numbers[1]


def load_questionnaire_catalog(doc_path: Path) -> dict[int, QuestionMeta]:
    lines = doc_path.read_text(encoding="utf-8").splitlines()
    questions: dict[int, QuestionMeta] = {}
    current_section = ""
    current_question: QuestionMeta | None = None
    option_mode = False

    for raw_line in lines:
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("## ") and not stripped.startswith("### "):
            current_section = stripped[3:].strip()
            current_question = None
            option_mode = False
            continue

        heading_match = QUESTION_HEADING_RE.match(stripped)
        if heading_match:
            number = int(heading_match.group(1))
            current_question = QuestionMeta(
                number=number,
                title=heading_match.group(2).strip(),
                section=current_section,
            )
            questions[number] = current_question
            option_mode = False
            continue

        if current_question is None:
            continue

        if not stripped:
            option_mode = False
            continue

        if stripped.startswith("- 题型："):
            current_question.question_type = stripped.split("：", 1)[1].strip()
            option_mode = False
            continue

        if stripped.startswith("- 选择限制："):
            min_selections, max_selections = parse_choice_limit(stripped.split("：", 1)[1].strip())
            current_question.min_selections = min_selections
            current_question.max_selections = max_selections
            option_mode = False
            continue

        if stripped.startswith("- 左锚点："):
            current_question.left_label = stripped.split("：", 1)[1].strip()
            option_mode = False
            continue

        if stripped.startswith("- 右锚点："):
            current_question.right_label = stripped.split("：", 1)[1].strip()
            option_mode = False
            continue

        if stripped == "- 选项：":
            option_mode = True
            continue

        if stripped.startswith("- ") and option_mode:
            option_text = stripped[2:].strip()
            if option_text:
                current_question.options.append(option_text)
            continue

        if stripped.startswith("- ") and current_question.question_type in {"单选", "单选题", "多选", "多选题"}:
            option_text = stripped[2:].strip()
            if option_text and "：" not in option_text:
                current_question.options.append(option_text)
                continue

        option_mode = False

    return questions


QUESTION_DOC_PATH = Path(__file__).resolve().parent / "docs" / "thu_questionnaire_template.md"
QUESTION_CATALOG = load_questionnaire_catalog(QUESTION_DOC_PATH)


def parse_args() -> argparse.Namespace:
    data_dir = Path(__file__).resolve().parent / "data"
    default_mode = os.getenv("MATCH_EXPLANATION_MODE", "llm")
    if default_mode not in {"llm", "heuristic"}:
        default_mode = "llm"
    provider = os.getenv("SOULAGENT_PROVIDER", "").strip().lower()
    default_api_base = os.getenv("SOULAGENT_API_BASE")
    if not default_api_base and provider == "siliconflow":
        default_api_base = "https://api.siliconflow.cn/v1"
    if not default_api_base:
        default_api_base = os.getenv("OPENAI_API_BASE")
    if not default_api_base:
        default_api_base = "https://api.openai.com/v1"
    default_model = os.getenv("SOULAGENT_MODEL", os.getenv("OPENAI_MODEL", "gpt-4.1-mini"))
    parser = argparse.ArgumentParser(description="Generate match explanation file via LLM or heuristic mode.")
    parser.add_argument("--matches", default=str(data_dir / "match_results.csv"), help="Path to match results CSV")
    parser.add_argument("--profiles", default=str(data_dir / "export.csv"), help="Path to questionnaire export CSV")
    parser.add_argument(
        "--output",
        default=str(data_dir / "match_explanations.csv"),
        help="Path to output file (.csv, .json, .jsonl, .xlsx)",
    )
    parser.add_argument("--mode", choices=["llm", "heuristic"], default=default_mode, help="Use LLM or local heuristic generator")
    parser.add_argument("--model", default=default_model, help="LLM model name for chat completions")
    parser.add_argument("--api-base", default=default_api_base, help="OpenAI-compatible API base URL")
    parser.add_argument("--api-key", default=None, help="API key. If omitted, read from environment variable")
    parser.add_argument("--api-key-env", default="OPENAI_API_KEY", help="Environment variable name for the API key")
    parser.add_argument("--timeout", type=int, default=120, help="Per-request timeout in seconds")
    parser.add_argument("--limit", type=int, default=None, help="Only process the first N matches")
    parser.add_argument("--sleep-seconds", type=float, default=0.0, help="Sleep between API requests")
    parser.add_argument("--include-raw-json", action="store_true", help="Include raw agent JSON in the output")
    parser.add_argument("--workers", type=int, default=8, help="Number of concurrent threads for LLM requests")
    return parser.parse_args()


def format_value(value: Any) -> str:
    if value is None:
        return "未作答"
    if isinstance(value, set):
        return "；".join(sorted(str(item) for item in value)) if value else "未作答"
    if isinstance(value, dict):
        if value.get("unlimited"):
            return "不限"
        min_value = value.get("min")
        max_value = value.get("max")
        if min_value is None and max_value is None:
            return "未作答"
        if min_value is None:
            return f"<= {max_value}"
        if max_value is None:
            return f">= {min_value}"
        if min_value == 0 and max_value == 999:
            return "不限"
        return f"{min_value}-{max_value}"
    if isinstance(value, float):
        if math.isnan(value):
            return "未作答"
        return f"{value:.2f}"
    return str(value)


def get_question_group(q_num: int) -> str:
    for group_name, q_range in QUESTION_GROUPS:
        if q_num in q_range:
            return group_name
    return "其他"


def format_answer_line(user: User, q_num: int) -> str | None:
    value = user.q(q_num)
    if value is None:
        return None
    importance = " [高优先级]" if user.imp(q_num) else ""
    question = QUESTION_CATALOG.get(q_num)
    title = question.title if question else f"Q{q_num}"
    return f"- {title}：{format_value(value)}{importance}"


def build_user_profile(user: User) -> str:
    sections: list[str] = []
    for group_name, q_range in QUESTION_GROUPS:
        lines = []
        for q_num in PROFILE_QUESTIONS:
            if q_num not in q_range:
                continue
            line = format_answer_line(user, q_num)
            if line:
                lines.append(line)
        if lines:
            sections.append(f"{group_name}\n" + "\n".join(lines))
    return "\n\n".join(sections)


def as_set(value: Any) -> set[str]:
    if isinstance(value, set):
        return {str(item) for item in value if str(item).strip()}
    return set()


def safe_scalar_distance(a: Any, b: Any, scale: float = 6.0) -> float | None:
    if a is None or b is None:
        return None
    if not isinstance(a, (int, float)) or not isinstance(b, (int, float)):
        return None
    return abs(float(a) - float(b)) / scale


def list_to_text(items: list[str]) -> str:
    return "；".join(items) if items else ""


def sanitize_cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ")
    return " ".join(text.split())


def flatten_text_items(items: list[str]) -> str:
    cleaned = [sanitize_cell_text(item) for item in items if sanitize_cell_text(item)]
    return " ｜ ".join(cleaned)


def build_export_explanation_text(agent_result: dict[str, Any]) -> str:
    parts: list[str] = []

    summary = sanitize_cell_text(agent_result.get("match_summary", ""))
    if summary:
        parts.append(summary)

    reasons = flatten_text_items(agent_result.get("match_reasons", []))
    if reasons:
        parts.append(f"匹配原因：{reasons}")

    topics = flatten_text_items(agent_result.get("date_topics", []))
    if topics:
        parts.append(f"推荐话题：{topics}")

    cautions = flatten_text_items(agent_result.get("cautions", []))
    if cautions:
        parts.append(f"需要留意：{cautions}")

    return " ".join(parts)


def save_output_file(output_df: pd.DataFrame, output_path: Path) -> None:
    suffix = output_path.suffix.lower()

    if suffix == ".xlsx":
        output_df.to_excel(output_path, index=False, engine="openpyxl")
        autosize_excel_columns(output_path)
        return

    if suffix == ".json":
        output_path.write_text(
            json.dumps(output_df.to_dict(orient="records"), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return

    if suffix == ".jsonl":
        output_df.to_json(output_path, orient="records", lines=True, force_ascii=False)
        return

    output_df.to_csv(output_path, index=False, encoding="utf-8-sig")


def build_tabular_output_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    tabular_rows: list[dict[str, Any]] = []
    for row in rows:
        explanation = row.get("explanation")
        tabular_row = {
            "match_group_no": row.get("match_group_no"),
            "group_name": row.get("group_name"),
            "explanation_text": row.get("explanation_text", ""),
            "explanation_json": json.dumps(explanation, ensure_ascii=False) if explanation is not None else "",
        }
        tabular_rows.append(tabular_row)
    return tabular_rows


def intersection_text(user_a: User, user_b: User, q_num: int) -> list[str]:
    return sorted(as_set(user_a.q(q_num)) & as_set(user_b.q(q_num)))


def mutual_trait_fit(user_a: User, user_b: User) -> tuple[list[str], list[str]]:
    a_wants = as_set(user_a.q(17))
    b_offers = as_set(user_b.q(18))
    b_wants = as_set(user_b.q(17))
    a_offers = as_set(user_a.q(18))
    return sorted(a_wants & b_offers), sorted(b_wants & a_offers)


def build_pair_evidence(user_a: User, user_b: User, match_row: pd.Series) -> dict[str, Any]:
    a_breakdown = compute_score_breakdown(user_a, user_b)
    b_breakdown = compute_score_breakdown(user_b, user_a)

    averaged_dimensions = {
        key: round((a_breakdown["dimensions"][key] + b_breakdown["dimensions"][key]) / 2, 4)
        for key in a_breakdown["dimensions"]
    }
    top_dimensions = sorted(averaged_dimensions.items(), key=lambda item: item[1], reverse=True)[:3]
    weak_dimensions = sorted(averaged_dimensions.items(), key=lambda item: item[1])[:2]
    trait_fit_a, trait_fit_b = mutual_trait_fit(user_a, user_b)

    q47_gap = safe_scalar_distance(user_a.q(47), user_b.q(47))
    q48_gap = safe_scalar_distance(user_a.q(48), user_b.q(48))

    evidence = {
        "scores": {
            "score_a_to_b": float(match_row["score_a_to_b"]),
            "score_b_to_a": float(match_row["score_b_to_a"]),
            "match_quality": float(match_row["match_quality"]),
            "grade": str(match_row["grade"]),
            "a_breakdown": a_breakdown,
            "b_breakdown": b_breakdown,
            "averaged_dimensions": averaged_dimensions,
            "top_dimensions": [
                {"label": DIMENSION_LABELS[key], "score": score} for key, score in top_dimensions
            ],
            "weak_dimensions": [
                {"label": DIMENSION_LABELS[key], "score": score} for key, score in weak_dimensions
            ],
        },
        "shared_points": {
            "shared_interests": intersection_text(user_a, user_b, 32),
            "shared_date_activities": intersection_text(user_a, user_b, 23),
            "shared_heartbeats": intersection_text(user_a, user_b, 20),
            "shared_dealbreakers": intersection_text(user_a, user_b, 52),
            "mutual_traits_for_a": trait_fit_a,
            "mutual_traits_for_b": trait_fit_b,
        },
        "attention_points": {
            "marriage_gap": q47_gap,
            "children_gap": q48_gap,
            "different_future_city": user_a.q(11) != user_b.q(11) if user_a.q(11) and user_b.q(11) else None,
            "different_relation_model": user_a.q(50) != user_b.q(50) if user_a.q(50) and user_b.q(50) else None,
        },
    }
    return evidence


def build_prompt(user_a: User, user_b: User, match_row: pd.Series, evidence: dict[str, Any]) -> tuple[str, str]:
    system_prompt = (
        "你是一个严谨的匹配结果解释 Agent，同时也是一个贴心的情感顾问，你现在要根据两个匹配对象的问卷答案和评分证据进行分析，给出两人初步相处的建议。"
        "这两个人未曾相识，你需要非常委婉或者间接地指出摩擦点以及婚姻、生育、性生活等敏感话题。"
        "你只能基于给定的问卷答案和评分证据做分析，不要臆测隐私，不要夸大结论。"
        "输出必须是 JSON 对象，不能包含 Markdown 代码块。"
        "**不要提及题号或任意一方的名字**，用自然中文总结。"
        "对于某些社会观念上不太好的习惯或品质，你应该用比较委婉的方式表达，比如“情感需求较强”而不是“这个人很啰嗦”之类的措辞。"
    )

    user_prompt = f"""
请分析这对匹配对象为什么可能适合，并推荐首次约会聊天话题。

要求：
1. 可以引用双方答案里能支撑结论的信息，但不要提及题号或任意一方的名字或昵称，比如直接说“你们都善于提供情绪价值”，而不是分别说两个人的性格；**一定要复合初见的场景**，例如“一起探讨婚姻生育观念”、“对于发生性关系的期待有差异”就是非常不合适的，你可以委婉表述为“对于长期关系的理念有差异”。
2. 区分“明显共同点”“互补点”“需要留意的点”；其中“需要留意的点”一定要委婉、概括性地表达，不宜提及具体的选项差异和敏感话题。
3. 约会话题要自然、具体，优先从双方共同兴趣、活动偏好、长期规划切入。
4. 语气真诚克制，避免空话。
5. 输出 JSON，格式如下：
{{
  "match_summary": "50-90字总结",
  "match_reasons": ["原因1", "原因2", "原因3"],
  "date_topics": ["话题1", "话题2", "话题3", "话题4", "话题5"],
  "cautions": ["留意点1", "留意点2"]
}}

匹配基础信息：
- 配对池：{match_row["pool"]}
- A：{user_a.nickname} (ID={user_a.user_id}, {user_a.department}, {user_a.grade})
- B：{user_b.nickname} (ID={user_b.user_id}, {user_b.department}, {user_b.grade})
- A->B 得分：{float(match_row["score_a_to_b"]):.4f}
- B->A 得分：{float(match_row["score_b_to_a"]):.4f}
- 综合匹配度：{float(match_row["match_quality"]):.4f}
- 等级：{match_row["grade"]}

维度证据：
{json.dumps(evidence, ensure_ascii=False, indent=2)}

A 的答案摘要：
{build_user_profile(user_a)}

B 的答案摘要：
{build_user_profile(user_b)}
""".strip()

    return system_prompt, user_prompt


def extract_json_from_text(content: str) -> dict[str, Any]:
    content = content.strip()
    if content.startswith("```"):
        lines = content.splitlines()
        if len(lines) >= 3:
            content = "\n".join(lines[1:-1]).strip()
    try:
        return json.loads(content)
    except json.JSONDecodeError:
        start = content.find("{")
        end = content.rfind("}")
        if start >= 0 and end > start:
            return json.loads(content[start:end + 1])
        raise


def call_chat_completions(
    client: OpenAI,
    model: str,
    system_prompt: str,
    user_prompt: str,
) -> dict[str, Any]:
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_prompt},
    ]

    try:
        response = client.chat.completions.create(
            model=model,
            temperature=0.6,
            response_format={"type": "json_object"},
            messages=messages,
        )
    except Exception:
        response = client.chat.completions.create(
            model=model,
            temperature=0.6,
            messages=messages,
        )

    message = response.choices[0].message.content
    if isinstance(message, list):
        text_parts = []
        for item in message:
            if isinstance(item, dict) and item.get("type") == "text":
                text_parts.append(item.get("text", ""))
        message = "\n".join(text_parts)
    if not isinstance(message, str):
        raise ValueError("Unexpected response content format from LLM API.")
    return extract_json_from_text(message)


def heuristic_reason_from_dimension(label: str) -> str:
    templates = {
        "匹配目的与阶段": "两人在参与目的、关系阶段和现实匹配条件上比较接近，推进节奏更容易对齐。",
        "地域与未来规划": "两人对未来落点和异地承受度相对接近，现实推进成本较低。",
        "吸引力与性格互补": "双方对欣赏的特质和自己能提供的特质有明显互相命中，容易产生吸引。",
        "生活方式": "作息、见面频率、个人空间和日常习惯相容度较高，长期相处摩擦可能更少。",
        "兴趣重合": "共同兴趣和活动偏好比较多，前期破冰和安排见面都更自然。",
        "安全感与亲密关系": "双方对安全感、沟通和亲密关系推进的期待较接近，更容易建立稳定连接。",
        "长期关系与未来观": "在关系模式、重大决定方式和长期关系观上更容易形成共识。",
    }
    return templates.get(label, f"在“{label}”这一维度上，两人的兼容度较高。")


def heuristic_topic_candidates(evidence: dict[str, Any], user_a: User, user_b: User) -> list[str]:
    topics: list[str] = []
    shared_interests = evidence["shared_points"]["shared_interests"]
    shared_activities = evidence["shared_points"]["shared_date_activities"]

    for item in shared_interests[:3]:
        topics.append(f"从共同兴趣“{item}”聊起，顺便交换各自最近的体验和偏好。")
    for item in shared_activities[:2]:
        topics.append(f"围绕两人都偏好的“{item}”讨论第一次见面会不会合适。")

    if user_a.q(11) and user_b.q(11):
        topics.append(f"聊聊彼此对未来城市规划的想法，以及希望怎样安排学习、工作和关系。")
    if user_a.q(20) or user_b.q(20):
        topics.append("分享各自最容易心动的相处方式，看看怎样的互动会让彼此舒服。")
    if user_a.q(23) or user_b.q(23):
        topics.append("聊第一次见面更偏向轻松散步、咖啡馆，还是一起做一件具体的事。")

    deduped = []
    seen = set()
    for item in topics:
        if item not in seen:
            deduped.append(item)
            seen.add(item)
    return deduped[:5]


def build_heuristic_result(user_a: User, user_b: User, evidence: dict[str, Any]) -> dict[str, Any]:
    top_dimensions = evidence["scores"]["top_dimensions"]
    weak_dimensions = evidence["scores"]["weak_dimensions"]
    shared_interests = evidence["shared_points"]["shared_interests"]
    shared_activities = evidence["shared_points"]["shared_date_activities"]
    shared_heartbeats = evidence["shared_points"]["shared_heartbeats"]
    mutual_traits_for_a = evidence["shared_points"]["mutual_traits_for_a"]
    mutual_traits_for_b = evidence["shared_points"]["mutual_traits_for_b"]

    summary_parts = []
    if shared_interests:
        summary_parts.append("兴趣上有明显重叠")
    if mutual_traits_for_a or mutual_traits_for_b:
        summary_parts.append("双方对彼此欣赏的特质有互相命中")
    if top_dimensions:
        summary_parts.append(f"在{top_dimensions[0]['label']}维度表现尤其稳定")
    if not summary_parts:
        summary_parts.append("整体匹配度较均衡")

    reasons = []
    for item in top_dimensions:
        reasons.append(heuristic_reason_from_dimension(item["label"]))
    if mutual_traits_for_a or mutual_traits_for_b:
        reasons.append(
            f"A 看重的特质里，B 命中了 {list_to_text(mutual_traits_for_a[:3]) or '部分关键项'}；"
            f"B 看重的特质里，A 命中了 {list_to_text(mutual_traits_for_b[:3]) or '部分关键项'}。"
        )
    if shared_interests or shared_activities or shared_heartbeats:
        reasons.append(
            "双方在兴趣、约会活动或容易心动的场景上存在交集，前期建立熟悉感会更快。"
        )

    cautions = []
    for item in weak_dimensions:
        cautions.append(f"“{item['label']}”相对不是最强项，前期沟通时可以多确认彼此预期。")
    if evidence["attention_points"]["different_future_city"]:
        cautions.append("双方对未来城市规划不完全一致，关系推进后需要更早讨论现实安排。")
    if evidence["attention_points"]["different_relation_model"]:
        cautions.append("双方理想关系模式不完全一致，适合在熟悉后尽早对齐边界和期待。")
    if not cautions:
        cautions.append("整体没有明显硬冲突，但仍建议在相处中尽早确认节奏和边界。")

    return {
        "match_summary": "；".join(summary_parts[:3]) + "，因此这对匹配更可能在相处初期快速建立连接。",
        "match_reasons": reasons[:3],
        "date_topics": heuristic_topic_candidates(evidence, user_a, user_b),
        "cautions": cautions[:2],
    }


def normalize_agent_result(result: dict[str, Any]) -> dict[str, Any]:
    return {
        "match_summary": str(result.get("match_summary", "")).strip(),
        "match_reasons": [str(item).strip() for item in result.get("match_reasons", []) if str(item).strip()][:3],
        "date_topics": [str(item).strip() for item in result.get("date_topics", []) if str(item).strip()][:5],
        "cautions": [str(item).strip() for item in result.get("cautions", []) if str(item).strip()][:2],
    }


def print_pair_result(
    user_a: User,
    user_b: User,
    match_row: pd.Series,
    agent_result: dict[str, Any],
) -> None:
    print("=" * 80)
    print(
        f"{user_a.nickname} (ID={user_a.user_id}) <-> "
        f"{user_b.nickname} (ID={user_b.user_id}) | "
        f"pool={match_row['pool']} | quality={float(match_row['match_quality']):.4f} | grade={match_row['grade']}"
    )
    print(f"匹配总结：{agent_result['match_summary']}")
    print("匹配原因：")
    for item in agent_result["match_reasons"]:
        print(f"- {item}")
    print("推荐话题：")
    for item in agent_result["date_topics"]:
        print(f"- {item}")
    print("需要留意：")
    for item in agent_result["cautions"]:
        print(f"- {item}")


def generate_result_for_pair(
    user_a: User,
    user_b: User,
    match_row: pd.Series,
    args: argparse.Namespace,
    api_key: str | None,
    client: OpenAI | None,
) -> tuple[dict[str, Any], dict[str, Any], str]:
    evidence = build_pair_evidence(user_a, user_b, match_row)

    if args.mode == "heuristic":
        return normalize_agent_result(build_heuristic_result(user_a, user_b, evidence)), evidence, "heuristic"

    if not api_key:
        raise ValueError(
            f"LLM mode requires an API key. Pass --api-key or set {args.api_key_env}."
        )
    if client is None:
        raise ValueError("LLM client was not initialized.")

    system_prompt, user_prompt = build_prompt(user_a, user_b, match_row, evidence)
    agent_result = call_chat_completions(
        client=client,
        model=args.model,
        system_prompt=system_prompt,
        user_prompt=user_prompt,
    )
    return normalize_agent_result(agent_result), evidence, "llm"


def autosize_excel_columns(output_path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(output_path)
    worksheet = workbook.active
    for column_cells in worksheet.columns:
        max_length = 0
        column_index = column_cells[0].column
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 80))
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(max_length + 2, 60)
    workbook.save(output_path)


# def build_prompt(user_a: User, user_b: User, match_row: pd.Series, evidence: dict[str, Any]) -> tuple[str, str]:
#     system_prompt = (
#         "你是一个严谨的匹配解释 Agent，同时是一个情感专家，你将要根据两个用户的生活习惯、兴趣爱好、价值观念等信息进行分析，给出相处建议。"
#         "你只能基于给定的问卷答案和评分证据分析，不要编造信息。"
#         "输出必须是 JSON 对象，不要输出 Markdown，不要解释格式。"
#         "语言自然、具体、克制。"
#     )

#     user_prompt = f"""
# 请为这对匹配对象生成三个子内容：

# 1. `match_reasons`：给出 2 点匹配原因，每一点都尽量写得充实、有内容，80-140 字，必须能落到双方答案或评分证据上。
# 1. `match_reasons`：给出 2 点匹配原因，每一点都尽量详细、细致、具体，建议写到 100-180 字，必须能落到问卷答案或评分证据上，尽量把“为什么适合”“适合在什么地方”“这种契合会如何体现在相处里”写清楚。
# 2. `date_topics`：给出 2 到 4 个约会话题，每一点都必须是很短的短语，不要写完整句子，长度尽量控制在 6-16 个字，风格参考“周边的辣味餐厅”“一起约球或看球赛”“各自的编程项目经历”。
# 3. `cautions`：给出 1 点需要留意，只写一句话，简洁直接。

# 要求：
# - 不要提题号。
# - 不要空泛夸赞，要写出为什么匹配。
# - 匹配原因里不要使用“双方”“他们”等指代，统一用“你们”或者直接用 `{user_a.nickname}`、`{user_b.nickname}`。
# - `cautions` 里也不要使用 A、B、双方、他们，统一用“你们”或者直接用 `{user_a.nickname}`、`{user_b.nickname}`。
# - 匹配原因优先写“共同点”和“互相命中”的地方。
# - 话题推荐尽量具体，像真实的人会抛出的开场，不能推荐过于露骨的话题，例如“一起讨论婚姻和孩子的想法”就是不适合初见双方讨论的话题。
# - 需要留意优先写真实但不夸大的潜在摩擦点，不要给出具体的选项，也不要直呼任意一方的名字，一定要用委婉的方式提醒双方。
# - 严格按下面 JSON 输出：
# {{
#   "match_reasons": ["原因1", "原因2"],
#   "date_topics": ["话题1", "话题2"],
#   "cautions": ["一句话留意点"]
# }}

# 匹配基础信息：
# - 配对池：{match_row["pool"]}
# - A：{user_a.nickname} (ID={user_a.user_id}, {user_a.department}, {user_a.grade})
# - B：{user_b.nickname} (ID={user_b.user_id}, {user_b.department}, {user_b.grade})
# - A->B 得分：{float(match_row["score_a_to_b"]):.4f}
# - B->A 得分：{float(match_row["score_b_to_a"]):.4f}
# - 综合匹配度：{float(match_row["match_quality"]):.4f}
# - 等级：{match_row["grade"]}

# 维度证据：
# {json.dumps(evidence, ensure_ascii=False, indent=2)}

# A 的答案摘要：
# {build_user_profile(user_a)}

# B 的答案摘要：
# {build_user_profile(user_b)}
# """.strip()

#     return system_prompt, user_prompt


def heuristic_topic_candidates(evidence: dict[str, Any], user_a: User, user_b: User) -> list[str]:
    topics: list[str] = []
    shared_interests = evidence["shared_points"]["shared_interests"]
    shared_activities = evidence["shared_points"]["shared_date_activities"]
    shared_heartbeats = evidence["shared_points"]["shared_heartbeats"]

    for item in shared_interests[:2]:
        topics.append(f"{item}相关经历")
    for item in shared_activities[:1]:
        topics.append(f"{item}第一次安排")
    for item in shared_heartbeats[:1]:
        topics.append(f"{item}的加分细节")

    if len(topics) < 4 and user_a.q(11) and user_b.q(11):
        topics.append("未来城市生活想象")
    if len(topics) < 4 and (user_a.q(32) or user_b.q(32)):
        topics.append("最想安利的兴趣坑")
    if len(topics) < 4:
        topics.append("最加分的相处细节")

    deduped = []
    seen = set()
    for item in topics:
        if item not in seen:
            deduped.append(item)
            seen.add(item)
    return deduped[:4]


def build_heuristic_result(user_a: User, user_b: User, evidence: dict[str, Any]) -> dict[str, Any]:
    top_dimensions = evidence["scores"]["top_dimensions"]
    weak_dimensions = evidence["scores"]["weak_dimensions"]
    shared_interests = evidence["shared_points"]["shared_interests"]
    shared_activities = evidence["shared_points"]["shared_date_activities"]
    shared_heartbeats = evidence["shared_points"]["shared_heartbeats"]
    mutual_traits_for_a = evidence["shared_points"]["mutual_traits_for_a"]
    mutual_traits_for_b = evidence["shared_points"]["mutual_traits_for_b"]

    reasons: list[str] = []
    if top_dimensions:
        reasons.append(
            f"你们在“{top_dimensions[0]['label']}”上表现最稳，说明核心关系预期更容易对齐。"
            f"这组匹配不是停留在泛泛的好感上，而是落在相处节奏、现实推进条件或长期看法这些真正影响关系质量的地方，所以会显得更扎实。"
        )
    if mutual_traits_for_a or mutual_traits_for_b:
        reasons.append(
            f"{user_a.nickname}看重的特质里，{user_b.nickname}命中了 {list_to_text(mutual_traits_for_a[:3]) or '若干关键项'}；"
            f"{user_b.nickname}看重的特质里，{user_a.nickname}也命中了 {list_to_text(mutual_traits_for_b[:3]) or '若干关键项'}。"
            f"这种互相命中的感觉，比单方面觉得“还不错”更容易在真实聊天和见面里持续升温。"
        )
    if len(reasons) < 2 and (shared_interests or shared_activities or shared_heartbeats):
        reasons.append(
            f"你们在共同兴趣、约会活动或心动场景上存在明显交集，前期不太需要靠硬聊来破冰。"
            f"像 {list_to_text((shared_interests + shared_activities + shared_heartbeats)[:3]) or '这些具体偏好'} 这类重合点，会让第一次见面更容易进入自然、有来回的状态。"
        )
    if len(reasons) < 2 and len(top_dimensions) > 1:
        reasons.append(
            f"除了最强项之外，“{top_dimensions[1]['label']}”也处在较高水平，这说明你们不是单点高分，而是多个维度同时不差。"
            f"这种更均衡的匹配，通常比只靠某一个兴趣点撑起来的关系更稳一些。"
        )

    cautions: list[str] = []
    for item in weak_dimensions:
        cautions.append(f"你们在“{item['label']}”上不是最强项，熟悉后最好尽早把预期说开。")
    if evidence["attention_points"]["different_future_city"]:
        cautions.append(f"{user_a.nickname}和{user_b.nickname}对未来城市规划不完全一致，越往后越需要正面聊现实安排。")
    if evidence["attention_points"]["different_relation_model"]:
        cautions.append(f"{user_a.nickname}和{user_b.nickname}理想中的关系模式不完全一样，边界和节奏最好早点确认。")
    if not cautions:
        cautions.append("你们整体没有明显硬冲突，但相处后仍建议尽早确认节奏与边界。")

    reasons = reasons[:2]
    topics = heuristic_topic_candidates(evidence, user_a, user_b)
    caution = cautions[:1]
    return {
        "match_summary": " ".join(reasons),
        "match_reasons": reasons,
        "date_topics": topics,
        "cautions": caution,
    }


def normalize_agent_result(result: dict[str, Any]) -> dict[str, Any]:
    reasons = [str(item).strip() for item in result.get("match_reasons", []) if str(item).strip()][:2]
    topics = [str(item).strip() for item in result.get("date_topics", []) if str(item).strip()][:4]
    if len(topics) < 2:
        fallback_topics = [item for item in result.get("date_topics", []) if str(item).strip()]
        topics = [str(item).strip() for item in fallback_topics][:2]
    cautions = [str(item).strip() for item in result.get("cautions", []) if str(item).strip()][:1]
    return {
        "match_summary": " ".join(reasons),
        "match_reasons": reasons,
        "date_topics": topics,
        "cautions": cautions,
    }


def print_pair_result(
    user_a: User,
    user_b: User,
    match_row: pd.Series,
    agent_result: dict[str, Any],
) -> None:
    print("=" * 80)
    print(
        f"{user_a.nickname} (ID={user_a.user_id}) <-> "
        f"{user_b.nickname} (ID={user_b.user_id}) | "
        f"pool={match_row['pool']} | quality={float(match_row['match_quality']):.4f} | grade={match_row['grade']}"
    )
    print("匹配原因：")
    for item in agent_result["match_reasons"]:
        print(f"- {item}")
    print("话题推荐：")
    for item in agent_result["date_topics"]:
        print(f"- {item}")
    print("需要留意：")
    for item in agent_result["cautions"]:
        print(f"- {item}")


def main() -> None:
    args = parse_args()
    api_key = (
        args.api_key
        or os.getenv("SOULAGENT_API_KEY")
        or os.getenv(args.api_key_env)
    )
    client = None
    if args.mode == "llm":
        print(f"[match-explanations] provider={os.getenv('SOULAGENT_PROVIDER', '').strip() or 'openai-compatible'} model={args.model} base_url={args.api_base}")
        client = OpenAI(
            api_key=api_key,
            base_url=args.api_base,
            timeout=args.timeout,
        )

    users = load_users(args.profiles)
    users_by_id = {user.user_id: user for user in users}
    matches_df = pd.read_csv(args.matches, encoding="utf-8-sig")

    if args.limit is not None:
        matches_df = matches_df.head(args.limit)

    total = len(matches_df)
    print_lock = threading.Lock()
    completed_count = [0]

    def _process_one(index: int, match_row: pd.Series) -> dict[str, Any]:
        match_group_no = index + 1
        user_a_id = int(match_row["user_a_id"])
        user_b_id = int(match_row["user_b_id"])
        user_a = users_by_id.get(user_a_id)
        user_b = users_by_id.get(user_b_id)

        if user_a is None or user_b is None:
            return {
                "match_group_no": match_group_no,
                "group_name": sanitize_cell_text(match_row.get("pool", "")),
                "explanation_text": sanitize_cell_text(f"生成失败：未找到用户资料 A={user_a_id}, B={user_b_id}"),
                "explanation": {"error": f"未找到用户资料 A={user_a_id}, B={user_b_id}"},
            }

        try:
            agent_result, evidence, generation_mode = generate_result_for_pair(
                user_a, user_b, match_row, args, api_key, client,
            )
            with print_lock:
                completed_count[0] += 1
                print(f"[{completed_count[0]}/{total}] {user_a.nickname} <-> {user_b.nickname} ✓")
                print_pair_result(user_a, user_b, match_row, agent_result)
            if args.mode == "llm" and args.sleep_seconds > 0:
                time.sleep(args.sleep_seconds)
            return {
                "match_group_no": match_group_no,
                "group_name": sanitize_cell_text(match_row.get("pool", "")),
                "explanation_text": build_export_explanation_text(agent_result),
                "explanation": agent_result,
            }
        except Exception as exc:
            with print_lock:
                completed_count[0] += 1
                print(f"[{completed_count[0]}/{total}] {user_a.nickname} <-> {user_b.nickname} ✗ {exc}")
            return {
                "match_group_no": match_group_no,
                "group_name": sanitize_cell_text(match_row.get("pool", "")),
                "explanation_text": sanitize_cell_text(f"生成失败：{exc}"),
                "explanation": {"error": str(exc)},
            }

    workers = max(1, args.workers) if args.mode == "llm" else 1
    print(f"[match-explanations] Processing {total} pairs with {workers} worker(s)...")

    tasks = list(matches_df.iterrows())
    output_map: dict[int, dict[str, Any]] = {}

    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = {
            executor.submit(_process_one, idx, row): idx
            for idx, row in tasks
        }
        for future in as_completed(futures):
            idx = futures[future]
            output_map[idx] = future.result()

    output_rows = [output_map[idx] for idx, _ in tasks]

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.suffix.lower() in {".csv", ".xlsx"}:
        output_df = pd.DataFrame(build_tabular_output_rows(output_rows))
    else:
        output_df = pd.DataFrame(output_rows)
    save_output_file(output_df, output_path)
    print(f"Saved explanation file to: {output_path}")


if __name__ == "__main__":
    main()
